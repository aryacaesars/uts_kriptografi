"""Build the member-1 handover report from actual browser evidence and source.
Run from the project root: python -X utf8 docs/build_report.py
Requires python-docx. PDF rendering is performed by docs/render_report.mjs.
"""
from pathlib import Path
from html import escape
import json
from datetime import datetime
from docx import Document
from docx.shared import Inches, Pt, Cm

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs"
evidence = json.loads((OUT / "bukti/browser-results.json").read_text(encoding="utf-8"))
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Cm(2)
section.left_margin = section.right_margin = Cm(2)
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)
doc.styles["Normal"].paragraph_format.space_after = Pt(7)
parts, md = [], []

def heading(text, level=1):
    doc.add_heading(text, level)
    parts.append(f"<h{level}>{escape(text)}</h{level}>")
    md.append(f"{'#' * level} {text}\n")

def paragraph(text):
    doc.add_paragraph(text)
    parts.append(f"<p>{escape(text)}</p>")
    md.append(text + "\n")

def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = str(value)
    parts.append("<table><thead><tr>" + "".join(f"<th>{escape(h)}</th>" for h in headers) + "</tr></thead><tbody>" + "".join("<tr>" + "".join(f"<td>{escape(str(v))}</td>" for v in row) + "</tr>" for row in rows) + "</tbody></table>")
    md.extend(["| " + " | ".join(headers) + " |", "| " + " | ".join("---" for _ in headers) + " |"])
    md.extend("| " + " | ".join(str(v) for v in row) + " |" for row in rows)
    md.append("")

def newpage():
    doc.add_page_break()
    parts.append('<div class="pagebreak"></div>')

heading("Laporan Bagian Anggota 1", 1)
heading("KriptoLab — Project UTS Kriptografi", 2)
paragraph("Program Studi Informatika, Fakultas Teknik, Universitas Siliwangi. Dosen pengampu: Ir. Randi Rizal, Ph.D.")
paragraph("Lingkup: antarmuka web, sanitizer, Vigenere Standard, Auto-Key Vigenere, Affine Cipher, serta integrasi antarmuka dengan modul kelompok yang telah tersedia.")
paragraph("Identitas penyusun: [konfirmasi nama dan NPM Anggota 1 sebelum laporan digabungkan]. Kontrak resmi menempatkan Dhiya Ulhaq Prima Yuga (237006066) sebagai Anggota 1, sedangkan catatan teknis mencantumkan penyusun Arya Achmad Caesar. Paket ini mengikuti peran Anggota 1 sesuai instruksi pengguna.")
paragraph("Status dokumen: materi serah terima bagian Anggota 1, bukan laporan akhir seluruh kelompok. Hill telah diintegrasikan; QA lima media nyata dan kompilasi laporan kelompok masih menunggu.")
paragraph("Tanggal penyusunan: 8 Oktober 2026. Waktu bukti browser (UTC): " + evidence["runAt"] + ".")

heading("BAB I. Pendahuluan")
heading("1.1 Latar belakang", 2)
paragraph("Praktikum kriptografi memerlukan antarmuka yang menghubungkan konsep transformasi pesan dengan hasil yang dapat diperiksa. KriptoLab menyediakan ruang kerja untuk memilih cipher, memasukkan kunci, mengenkripsi atau mendekripsi pesan, dan meninjau representasi Base64. Pembedaan cipher alfabet dan cipher byte ditampilkan agar pengguna memahami data yang dipertahankan selama pemrosesan.")
heading("1.2 Tujuan", 2)
paragraph("Tujuan bagian Anggota 1 adalah mengimplementasikan tiga cipher klasik sesuai kontrak fungsi kelompok; menyediakan sanitasi alfabet A–Z; membangun antarmuka responsif untuk teks dan berkas; menghubungkan modul yang tersedia; serta menyiapkan bukti pengujian dan tangkapan layar untuk laporan kelompok.")
heading("1.3 Ruang lingkup dan pembagian", 2)
table(["Peran", "Lingkup"], [
    ["Anggota 1", "GUI, sanitizer, Vigenere Standard, Auto-Key, Affine, helper tampilan, materi UI/Bab I."],
    ["Anggota 2", "Playfair, Extended Vigenere, Super, Enigma, kontainer KRIP. Bahan tersedia dan diintegrasikan."],
    ["Anggota 3", "Hill 2x2/3x3, QA lima media nyata, kompilasi laporan akhir dan README kelompok."],
])
paragraph("Teks alfabet dipulihkan sebagai hasil sanitasi, sehingga spasi, angka, dan tanda baca tidak kembali. Pemulihan byte mentah beserta format file memakai Extended Vigenere atau Super Enkripsi. Implementasi menggunakan JavaScript/TypeScript sehingga bonus Ruby tidak diterapkan.")

heading("BAB II. Rancangan dan Algoritma Anggota 1")
heading("2.1 Arsitektur aplikasi", 2)
paragraph("Next.js 16.4.0 menyediakan halaman dan aset aplikasi. Komponen interaktif berjalan sebagai Client Component dengan React. Alur pemrosesan: form pengguna → Uint8Array melalui TextEncoder/File.arrayBuffer → processRequest → modul cipher → tampilan dan Blob unduhan. Pesan, file, serta kunci diproses di browser. Tidak ada endpoint server untuk mengirimkan data kriptografi.")
heading("2.2 Sanitasi dan representasi", 2)
paragraph("sanitizeText(text) membuang karakter selain a–z/A–Z sebelum mengubah menjadi huruf kapital. Contoh: Halo, Dunia 123! menjadi HALODUNIA. Penyaringan sebelum uppercase mencegah huruf non-ASCII yang berubah menjadi ASCII saat uppercase ikut masuk. Kunci Vigenere dan Auto-Key juga disanitasi; kunci yang tidak menyisakan huruf ditolak.")
paragraph("Base64 mengubah representasi byte menjadi teks dan tidak menambah kerahasiaan. Helper mengonversi string menjadi UTF-8 terlebih dahulu. Besar/kecil huruf Base64 dipertahankan; format huruf kapital/kecil hanya diterapkan pada keluaran cipher alfabet. Pratinjau layar dibatasi 12.000 karakter, tetapi salin dan unduh menggunakan hasil lengkap.")
heading("2.3 Vigenere Standard", 2)
paragraph("Huruf dipetakan A=0 hingga Z=25. Dengan panjang kunci m, enkripsi memakai C[i]=(P[i]+K[i mod m]) mod 26. Dekripsi memakai P[i]=(C[i]−K[i mod m]+26) mod 26. Kunci berulang sepanjang pesan. Contoh pemeriksaan independen: ATTACKATDAWN dengan LEMON menghasilkan LXFOPVEFRNHR.")
heading("2.4 Auto-Key Vigenere", 2)
paragraph("Aliran kunci dibentuk dari kunci awal diikuti plainteks, bukan pengulangan kunci. Enkripsi memakai C[i]=(P[i]+S[i]) mod 26. Dekripsi dilakukan berurutan: setelah posisi kunci awal terlewati, huruf plainteks yang sudah ditemukan menjadi kunci untuk posisi berikutnya. ATTACKATDAWN dengan QUEENLY menghasilkan QNXEPVYTWTWP. Memakai cipherteks sebagai tambahan aliran kunci akan menghasilkan varian yang berbeda dan tidak sesuai kontrak tugas ini.")
heading("2.5 Affine Cipher", 2)
paragraph("Enkripsi memakai C=(aP+b) mod 26, sedangkan dekripsi memakai P=a⁻¹(C−b) mod 26. Parameter a harus integer yang relatif prima dengan 26 dan b harus integer 0–25. Nilai a valid modulo 26 ialah 1,3,5,7,9,11,15,17,19,21,23,25. Invers dicari dengan menguji a·x mod 26=1. Untuk a=7 diperoleh invers 15. Input a=13 ditolak pada GUI dan fungsi karena tidak memiliki invers.")
heading("2.6 Integrasi berkas", 2)
paragraph("File dienkripsi sebagai byte mentah oleh Extended/Super, lalu dibungkus: 4 byte KRIP + 1 byte panjang ekstensi + ekstensi ASCII + payload terenkripsi. Header asli file berada di dalam payload dan ikut dienkripsi; KRIP merupakan header kontainer baru. unpackageBinaryFile mengembalikan originalExtension dan encryptedBytes sesuai kontrak resmi, kemudian payload didekripsi. Algoritma/kunci tidak disimpan dalam kontainer dan harus ditentukan pengguna.")
paragraph("Input dibatasi 10 MiB; unggahan dekripsi diperbolehkan memiliki tambahan maksimum 260 byte header kontainer agar batas tersebut tetap dapat dipulihkan. File alfabet hanya diproses sebagai teks UTF-8 dan hasil dekripsinya diunduh sebagai .txt. File tanpa ekstensi pada jalur byte menggunakan .bin.")

heading("BAB III. Antarmuka dan Cara Penggunaan")
paragraph("Antarmuka dibagi menjadi konfigurasi cipher, masukan, dan hasil. Form kunci berubah sesuai algoritma: teks untuk Vigenere/Auto-Key, pengali dan pergeseran untuk Affine, matriks untuk Hill, kunci transposisi untuk Super, serta rotor/ring/reflektor/plugboard untuk Enigma. Tema gelap/terang disimpan sebagai preferensi lokal. Layout dua kolom menjadi satu kolom pada layar kecil.")
paragraph("Alur teks: pilih cipher → tentukan kunci → ketik pesan → Enkripsi pesan → tinjau hasil → Balikkan proses → Dekripsi pesan. Alur file: pilih cipher byte → unggah file → enkripsi → Unduh hasil .dat → pilih Dekripsi dan unggah .dat → unduh file dengan ekstensi semula. Jika memakai hasil salinan, pilih encoding Base64 untuk payload byte. Metadata ekstensi hanya dipertahankan melalui kontainer .dat.")
paragraph("Hill menyediakan masukan matriks 2x2 atau 3x3 dengan validasi determinan langsung; tombol proses nonaktif jika matriks tidak valid. Padding X dipertahankan pada dekripsi. Pesan kesalahan ditampilkan dekat hasil dengan role alert; perubahan parameter membersihkan hasil lama untuk menghindari keluaran yang tidak sesuai konfigurasi saat ini.")

heading("BAB IV. Pengujian Bagian Anggota 1")
paragraph("Pengujian logika memakai test runner bawaan Node.js: 14 kelompok uji mencakup known-answer, round trip, sanitasi Unicode, kunci tidak valid, semua nilai a valid dan b 0–25, Base64, kontainer rusak, integrasi byte, batas ukuran 10 MiB, dan status Hill. Uji browser otomatis dijalankan pada Chromium " + evidence["browser"] + " terhadap build produksi lokal; 13 pemeriksaan tercatat tanpa page error, termasuk vektor Hill HELP -> HIAT dan round-trip di GUI.")
rows = []
names = {"vigenere": "Vigenere", "autokey": "Auto-Key", "affine": "Affine", "hill": "Hill"}
for item in evidence["checks"]:
    if isinstance(item, dict) and "ciphertext" in item:
        rows.append([names[item["algorithm"]], item["plaintext"], item["ciphertext"], "Sesuai"])
table(["Cipher", "Plainteks tersanitasi", "Cipherteks", "Dekripsi"], rows)
paragraph("Kunci contoh: Vigenere BAGUS; Auto-Key UNSIL; Affine a=7 dan b=10. File teks sintetis berisi Halo, Dunia! 123 diikuti baris Kriptografi. dipulihkan sebagai HALODUNIAKRIPTOGRAFI oleh ketiga cipher Anggota 1.")
paragraph("Pengujian browser juga mengunduh dan mengunggah kembali berkas sintetis 1.025 byte pada Extended dan Super. Byte hasil dibandingkan langsung dengan input, dan SHA-256 dicatat dalam browser-results.json. Hal ini membuktikan kesamaan byte untuk fixture yang diuji; tidak menggantikan bukti lima media asli, pengujian pemutar audio/video, atau penilaian keamanan algoritma.")
table(["Aspek", "Hasil / batas bukti"], [
    ["Vigenere, Auto-Key, Affine", "Known-answer, round trip teks, upload/download teks tersanitasi lulus."],
    ["Sanitizer / Base64", "Filter A–Z, UTF-8, byte 0–255, input tidak valid, clipboard browser lulus."],
    ["UI", "Tema gelap/terang, persistensi tema, error Affine, layar 390 px tanpa overflow horizontal lulus."],
    ["Integrasi Anggota 2", "Byte sintetis Extended/Super lulus; known-answer Playfair/Enigma lulus pada uji logika."],
    ["Hill", "Vektor 2x2/3x3, padding dan validasi matriks lulus."],
    ["QA lima media asli", "Belum dilakukan dalam paket Anggota 1; dilanjutkan Anggota 3."],
])
paragraph("Validasi build/lint dan log test disimpan di docs/bukti. Tabel evaluasi sembilan butir final wajib disusun setelah seluruh kelompok menyelesaikan integrasi serta QA; tabel rencana pada breakdown tidak disalin menjadi klaim keberhasilan.")

heading("BAB V. Kesimpulan dan Serah Terima")
paragraph("Bagian Anggota 1 menyediakan antarmuka yang dapat dijalankan, tiga cipher klasik sesuai kontrak, sanitasi, Base64, dan alur file terintegrasi. Pengujian yang dijalankan mendukung hasil pada kasus uji yang didokumentasikan. Keluaran alfabet tidak memulihkan karakter yang dibuang, sedangkan jalur byte memulihkan fixture yang diuji secara identik.")
paragraph("Anggota 3 perlu menguji dokumen/gambar/database/audio/video nyata, melengkapi tabel sembilan butir, dan menggabungkan laporan kelompok. Hill sudah tersedia di aplikasi. Anggota 2 perlu meninjau penjelasan teknis modul biner dan cipher miliknya. Identitas penyusun serta tenggat Classroom perlu dicocokkan sebelum penyerahan.")
heading("Acuan pekerjaan", 2)
for text in [
    "Soal UTS Kriptografi 20261.pdf: naskah lokal 3 halaman, spesifikasi program, GUI, file, dan laporan.",
    "breakdown tugas.txt; KONTRAK_INTERFACE_DAN_FUNGSI.txt; CATATAN_DAN_PANDUAN_ANGGOTA_1.txt; ALUR_KERJA_DAN_URUTAN_PENGERJAAN.txt: bahan lokal pembagian dan kontrak kelompok.",
    "Folder bahan ../src: sumber awal UI dan modul; disalin lalu disesuaikan di proyek uts-kriptografi. Atribusi modul Anggota 2 dipertahankan.",
    "Panduan lokal Next.js 16.4.0: node_modules/next/dist/docs/01-app/01-getting-started, bagian layouts/pages, server/client components, dan CSS.",
    "tests/ciphers.test.mjs; tests/browser-smoke.mjs; docs/bukti/browser-results.json: prosedur dan bukti pengujian aktual.",
]:
    paragraph(text)

screens = [
    ("01-vigenere-dark.png", "Vigenere Standard — tema gelap", "Pesan contoh KRIPTOGRAFI UNSIL, kunci BAGUS, dan hasil enkripsi yang dihitung aplikasi."),
    ("02-base64-light.png", "Representasi Base64 — tema terang", "Payload hasil serta pesan input ditampilkan dalam Base64. Salin hasil diuji melalui clipboard browser."),
    ("03-autokey-dekripsi.png", "Dekripsi Auto-Key", "Plainteks dipulihkan melalui tombol Balikkan proses dengan kunci awal UNSIL."),
    ("04-affine-validasi.png", "Validasi Affine", "Parameter a=13 tidak relatif prima dengan 26 sehingga pemrosesan ditolak dengan pesan kesalahan."),
    ("05-super-file.png", "Enkripsi dan unduhan .dat", "File sintetis 1.025 byte dienkripsi melalui Super. Ini bukti alur GUI, bukan media asli final kelompok."),
    ("06-super-file-dekripsi.png", "Pemulihan file dari .dat", "File .dat hasil unduhan diunggah kembali; hasil dekripsi dibandingkan byte demi byte dengan fixture asli."),
    ("07-mobile.png", "Antarmuka mobile 390 px", "Layout satu kolom; pengujian memeriksa tidak ada overflow horizontal pada viewport ini."),
    ("08-hill.png", "Hill Cipher dan validasi input", "Matriks 2×2 valid. Dekripsi ABC ditolak kerana panjang cipherteks bukan kelipatan ordo."),
]
for i, (filename, title, caption) in enumerate(screens, 1):
    newpage()
    heading(f"Lampiran A.{i}. {title}", 1)
    paragraph(caption)
    picture = OUT / "bukti" / filename
    from docx.image.image import Image
    img = Image.from_file(str(picture))
    width = min(6.3, 8.0 * img.px_width / img.px_height)
    doc.add_picture(str(picture), width=Inches(width))
    parts.append(f'<img class="screenshot" src="bukti/{filename}" alt="{escape(title)}">')
    md.append(f"![{title}](bukti/{filename})\n")

newpage()
heading("Lampiran B. Kode Modul Anggota 1")
paragraph("Kode di bawah diambil langsung dari implementasi yang diuji. Kode GUI dan integrasi lengkap tersedia di components/CipherWorkbench.jsx, app/globals.css, dan src/lib/workbench.js dalam paket proyek; modul anggota lain tetap berada di src/modules.")
for filename in ["sanitizer.js", "vigenere_standard.js", "autokey_vigenere.js", "affine_cipher.js", "display_helpers.js"]:
    heading(filename, 2)
    code = (ROOT / "src/modules" / filename).read_text(encoding="utf-8")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    parts.append(f"<pre>{escape(code)}</pre>")
    md.append("```js\n" + code + "```\n")

doc.core_properties.title = "Laporan Bagian Anggota 1 — KriptoLab"
doc.core_properties.subject = "Materi serah terima, bukan laporan final kelompok"
doc.core_properties.author = "Anggota 1 (identitas perlu dikonfirmasi)"
doc.save(OUT / "LAPORAN_ANGGOTA_1.docx")
(OUT / "LAPORAN_ANGGOTA_1.md").write_text("\n".join(md), encoding="utf-8")
css = """@page{size:A4;margin:18mm}body{font:11pt/1.5 Arial;color:#182b3a}h1{font-size:19pt;color:#125c4c;margin-top:20pt}h2{font-size:13pt;margin-top:16pt}h1,h2{break-after:avoid}p{text-align:justify}table{border-collapse:collapse;width:100%;font-size:9pt;table-layout:fixed;break-inside:avoid}th,td{border:1px solid #b9c9cc;padding:7px;overflow-wrap:anywhere;vertical-align:top}th{background:#e9f2ef;text-align:left}tr{break-inside:avoid}pre{break-inside:avoid;font:8pt/1.4 Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere;background:#f4f6f7;padding:10px}.pagebreak{break-before:page}.screenshot{display:block;max-width:100%;max-height:220mm;width:auto;height:auto;margin:10px auto}p,pre{orphans:3;widows:3}"""
(OUT / "LAPORAN_ANGGOTA_1.html").write_text('<!doctype html><html lang="id"><meta charset="utf-8"><title>Laporan Bagian Anggota 1</title><style>' + css + '</style><body>' + "\n".join(parts) + '</body></html>', encoding="utf-8")
print("Generated DOCX, Markdown, and HTML report from", evidence["runAt"])
